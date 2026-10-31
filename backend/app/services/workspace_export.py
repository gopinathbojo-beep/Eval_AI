"""Export only authorized exams; family exports never contain comparative grades."""

from io import BytesIO

from openpyxl import Workbook

from ..models import Evaluation, Exam
from ..workspaces import exams_query
from .school_reports import PASS_PERCENT, raw_percentage, relative_grades


def marksheet(db, user):
    wb = Workbook()
    ws = wb.active
    ws.title = "Evaluation results"
    columns = [
        "Student",
        "Class",
        "Roll number",
        "Exam",
        "Subject",
        "Academic year",
        "Marks",
        "Maximum",
        "Percentage",
        "Result",
    ]
    if user.role != "PARENT":
        columns.append("Grade")
    ws.append(columns)
    for exam in exams_query(db, user).order_by(Exam.created_at):
        evaluations = (
            db.query(Evaluation).filter_by(exam_id=exam.id, status="FINALIZED").all()
        )
        grades = relative_grades(
            {e.id: raw_percentage(e.total_marks, e.maximum_marks) for e in evaluations}
        )
        for ev in evaluations:
            pct = raw_percentage(ev.total_marks, ev.maximum_marks)
            row = [
                ev.student.name,
                exam.school_class.name,
                ev.student.roll_number,
                exam.name,
                exam.subject.name,
                exam.academic_year,
                ev.total_marks,
                ev.maximum_marks,
                round(pct, 2) if pct is not None else None,
                "Incomplete"
                if pct is None
                else "Pass"
                if pct >= PASS_PERCENT
                else "Fail",
            ]
            if user.role != "PARENT":
                row.append(grades[ev.id])
            # Names and other uploaded fields must remain text, never spreadsheet formulas.
            for col, value in enumerate(row, 1):
                cell = ws.cell(ws.max_row + 1 if col == 1 else ws.max_row, col, value)
                if isinstance(value, str):
                    cell.data_type = "s"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 22
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()
