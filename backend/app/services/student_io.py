"""Student list import (Excel) and export (Excel / PDF)."""
from __future__ import annotations
import re
from io import BytesIO
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session
from ..models import SchoolClass, Student

COLUMNS = ["Name", "Roll Number", "Registration Number", "Class", "Section"]
_HEADER_ALIASES = {
    "name": ("name", "student name", "studentname", "full name"),
    "roll_number": ("roll number", "roll no", "rollno", "roll", "roll num", "roll no."),
    "registration_number": ("registration number", "registration no", "reg no", "regno", "reg number",
                            "register number", "register no", "registration", "reg. no.", "reg no."),
    "class": ("class", "class name", "std", "standard", "grade"),
    "section": ("section", "sec", "division", "div"),
}


def _clean(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():  # Excel stores 12 as 12.0
        v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip()


def norm_reg(v: str | None) -> str:
    return re.sub(r"[^A-Z0-9]", "", (v or "").upper())


def class_key(name: str) -> str:
    """'10-A', '10 A' and '10a' are the same class."""
    return re.sub(r"[^a-z0-9]", "", name.lower())


def class_name_for(cls_name: str, section: str) -> str:
    """Each class + section is its own class: Class '10' and Section 'A' -> '10-A' (unless the name already has it)."""
    if not section or class_key(cls_name).endswith(class_key(section)):
        return cls_name
    return f"{cls_name}-{section}"


def split_class_name(name: str) -> tuple[str, str | None]:
    """'10-A' -> ('10', 'A'); a name without a section ('Grade 12') -> ('Grade 12', None)."""
    std, dash, sec = name.rpartition("-")
    return (std.strip(), sec.strip()) if dash and std.strip() and sec.strip() else (name.strip(), None)


def reg_taken(db: Session, teacher_id: str, reg: str | None, exclude_id: str | None = None) -> Student | None:
    """Registration numbers are unique across all of a teacher's students (they identify answer sheets)."""
    if not reg:
        return None
    key = norm_reg(reg)
    rows = (db.query(Student).join(SchoolClass, SchoolClass.id == Student.class_id)
            .filter(SchoolClass.teacher_id == teacher_id, Student.registration_number.isnot(None)).all())
    return next((s for s in rows if norm_reg(s.registration_number) == key and s.id != exclude_id), None)


def import_students(db: Session, teacher_id: str, data: bytes, default_class: SchoolClass | None) -> dict:
    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise ValueError("This file is not a readable Excel workbook (.xlsx).")
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    header_idx, cols, best = None, {}, {}
    for i, row in enumerate(rows[:10]):  # header may sit below a title row
        found = {}
        for j, cell in enumerate(row):
            h = _clean(cell).lower().rstrip(":")
            for field, names in _HEADER_ALIASES.items():
                if h in names and field not in found:
                    found[field] = j
        if "name" in found and "roll_number" in found:
            header_idx, cols = i, found
            break
        best = max(best, found, key=len)
    if header_idx is None:
        missing = " and ".join(f"'{label}'" for field, label in (("name", "Name"), ("roll_number", "Roll Number"))
                               if field not in best)
        raise ValueError(f"The sheet has no {missing} column. Add a {missing} column (roll numbers as students write them "
                         f"on their answer sheets) and import again.")
    if "class" not in cols and default_class is None:
        raise ValueError("The sheet has no 'Std' column. Add Std (and Sec) columns and import again.")

    classes = {class_key(c.name): c for c in db.query(SchoolClass).filter_by(teacher_id=teacher_id)}
    created = updated = 0
    errors: list[dict] = []
    seen_regs: dict[str, int] = {}
    seen_rolls: dict[tuple[str, str], int] = {}
    for n, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        get = lambda f: _clean(row[cols[f]]) if f in cols and cols[f] < len(row) else ""  # noqa: E731
        name, roll, reg, cls_name, section = get("name"), get("roll_number"), get("registration_number"), get("class"), get("section")
        if not any((name, roll, reg, cls_name, section)):
            continue
        if not name or not roll:
            errors.append({"row": n, "reason": "Name and roll number are required."})
            continue
        if len(name) > 200 or len(roll) > 50 or len(reg) > 50 or len(section) > 20 or len(cls_name) > 100:
            errors.append({"row": n, "reason": "A value is too long."})
            continue
        if cls_name:
            cls_name = class_name_for(cls_name, section)
            c = classes.get(class_key(cls_name))
            if c is None:
                c = SchoolClass(teacher_id=teacher_id, name=cls_name)
                db.add(c)
                db.flush()
                classes[class_key(cls_name)] = c
        elif default_class is not None:
            c = default_class
        else:
            errors.append({"row": n, "reason": "Std is required."})
            continue
        roll_key = (c.id, roll.lower())
        if roll_key in seen_rolls:
            errors.append({"row": n, "reason": f"Roll number {roll} appears twice for {c.name} (also row {seen_rolls[roll_key]})."})
            continue
        seen_rolls[roll_key] = n
        if reg:
            if norm_reg(reg) in seen_regs:
                errors.append({"row": n, "reason": f"Registration number {reg} appears twice (also row {seen_regs[norm_reg(reg)]})."})
                continue
            seen_regs[norm_reg(reg)] = n

        existing = reg_taken(db, teacher_id, reg) if reg else None
        by_roll = db.query(Student).filter_by(class_id=c.id, roll_number=roll).first()
        if existing and by_roll and existing.id != by_roll.id:
            errors.append({"row": n, "reason": f"Registration number {reg} belongs to {existing.name}, but roll number "
                                               f"{roll} in {c.name} is {by_roll.name}."})
            continue
        st = existing or by_roll
        if st is None:
            db.add(Student(class_id=c.id, name=name, roll_number=roll, registration_number=reg or None, section=section or None))
            created += 1
        else:
            if st.class_id != c.id and db.query(Student).filter_by(class_id=c.id, roll_number=roll).first():
                errors.append({"row": n, "reason": f"Roll number {roll} already exists in {c.name}."})
                continue
            st.class_id, st.name, st.roll_number = c.id, name, roll
            if reg:
                st.registration_number = reg
            if section:
                st.section = section
            updated += 1
        db.flush()
    return {"created": created, "updated": updated, "errors": errors}


def _rows(students: list[Student]) -> list[list[str]]:
    return [[s.name, s.roll_number, s.registration_number or "", s.school_class.name if s.school_class else "", s.section or ""]
            for s in students]


def export_xlsx(students: list[Student]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Students"
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="18447E")
    for r in _rows(students):
        ws.append(r)
    for col, width in zip("ABCDE", (32, 14, 22, 14, 10)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_pdf(students: list[Student], title: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=15 * mm, bottomMargin=15 * mm,
                            title=title)
    styles = getSampleStyleSheet()
    data = [["#", *COLUMNS]] + [[str(i), *r] for i, r in enumerate(_rows(students), 1)]
    t = Table(data, repeatRows=1, colWidths=[10 * mm, 60 * mm, 25 * mm, 40 * mm, 25 * mm, 20 * mm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#18447E")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    doc.build([Paragraph(title, styles["Title"]), Paragraph(f"{len(students)} students", styles["Normal"]), Spacer(1, 6 * mm), t])
    return buf.getvalue()
