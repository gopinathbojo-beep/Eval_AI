from datetime import datetime, timezone
from io import BytesIO
from uuid import uuid4

import fitz
from app.config import get_settings
from app.database import SessionLocal
from app.main import app
from app.models import (
    AnswerSheetDocument,
    Evaluation,
    Teacher,
)
from app.security import create_token
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from tests.test_registration import HEADERS, PROFILE, signup


def setup_school(browser):
    signup(browser)
    assert (
        browser.put("/api/auth/profile", headers=HEADERS, json=PROFILE).status_code
        == 200
    )
    cls = browser.post(
        "/api/classes", headers=HEADERS, json={"std": "10", "section": "A"}
    ).json()
    subject = browser.post(
        "/api/subjects", headers=HEADERS, json={"name": "English"}
    ).json()
    return cls, subject


def exam(browser, cls, subject):
    r = browser.post(
        "/api/exams",
        headers=HEADERS,
        json={
            "name": "Term exam",
            "class_id": cls["id"],
            "subject_id": subject["id"],
            "exam_type": "Term 1",
            "total_marks": 80,
        },
    )
    assert r.status_code == 201, r.text
    return r.json()


def test_teacher_assignment_management_and_workspace_isolation(client):
    with (
        TestClient(app) as school,
        TestClient(app) as other,
        TestClient(app) as teacher,
    ):
        cls, subject = setup_school(school)
        c2, s2 = setup_school(other)
        e2 = exam(other, c2, s2)
        email = f"teacher-{uuid4().hex}@example.com"
        created = school.post(
            "/api/management/teachers",
            headers=HEADERS,
            json={
                "name": "Teacher",
                "email": email,
                "password": "temporary-password-123",
                "assignments": [{"class_id": cls["id"], "subject_id": subject["id"]}],
            },
        )
        assert created.status_code == 201, created.text
        tid = created.json()["id"]
        assert (
            teacher.post(
                "/api/auth/login",
                headers=HEADERS,
                json={
                    "email": email,
                    "password": "temporary-password-123",
                    "browser_session": True,
                },
            ).status_code
            == 200
        )
        assert teacher.get("/api/exams").status_code == 403
        assert (
            teacher.post(
                "/api/auth/password",
                headers=HEADERS,
                json={
                    "current_password": "temporary-password-123",
                    "password": "changed-password-123",
                },
            ).status_code
            == 200
        )
        assert [c["id"] for c in teacher.get("/api/classes").json()] == [cls["id"]]
        assert (
            teacher.post(
                "/api/students",
                headers=HEADERS,
                json={"name": "No", "roll_number": "3", "class_id": cls["id"]},
            ).status_code
            == 403
        )
        own = exam(teacher, cls, subject)
        management_exam = exam(school, cls, subject)
        assert teacher.get(f"/api/exams/{management_exam['id']}").status_code == 404
        assert school.get(f"/api/exams/{own['id']}").status_code == 200
        assert teacher.get(f"/api/exams/{e2['id']}").status_code == 404
        assert school.get(f"/api/exams/{e2['id']}").status_code == 404
        assert (
            school.patch(
                f"/api/management/teachers/{tid}",
                headers=HEADERS,
                json={"assignments": []},
            ).status_code
            == 200
        )
        assert (
            teacher.get("/api/exams").status_code == 401
        )  # assignment change revokes existing sessions
        teacher.post(
            "/api/auth/login",
            headers=HEADERS,
            json={
                "email": email,
                "password": "changed-password-123",
                "browser_session": True,
            },
        )
        assert teacher.get(f"/api/exams/{own['id']}").status_code == 404
        assert teacher.get("/api/exams").json() == []
        assert (
            school.patch(
                f"/api/management/teachers/{tid}",
                headers=HEADERS,
                json={"assignments": [{"class_id": c2["id"], "subject_id": s2["id"]}]},
            ).status_code
            == 404
        )


def test_school_disable_revokes_members_and_admin_protected(client):
    with TestClient(app) as school:
        setup_school(school)
        me = school.get("/api/auth/me").json()
        with SessionLocal() as db:
            owner = db.query(Teacher).filter_by(email=get_settings().admin_email).one()
            headers = {
                "Authorization": f"Bearer {create_token(owner.id, owner.session_version)}"
            }
        assert (
            client.patch(
                f"/api/admin/workspaces/{me['workspace_id']}",
                headers=headers,
                json={"is_active": False},
            ).status_code
            == 200
        )
        assert school.get("/api/auth/me").status_code == 401
        assert (
            client.patch(
                f"/api/admin/workspaces/{me['workspace_id']}",
                headers=headers,
                json={"is_active": True},
            ).status_code
            == 200
        )
        assert school.get("/api/auth/me").status_code == 401
        assert (
            client.patch(
                f"/api/admin/users/{owner.id}",
                headers=headers,
                json={"is_active": False},
            ).status_code
            == 400
        )


def test_import_preview_conflicts_and_confirmation(client):
    with TestClient(app) as school:
        setup_school(school)
        wb = Workbook()
        ws = wb.active
        ws.append(["Name", "Roll Number", "Class", "Section"])
        ws.append(["Aditi", "1", "12", "B"])
        buf = BytesIO()
        wb.save(buf)
        files = {
            "file": (
                "students.xlsx",
                buf.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        }
        before = school.get("/api/students").json()
        r = school.post("/api/students/import", headers=HEADERS, files=files)
        assert (
            r.status_code == 200
            and r.json()["created"] == 1
            and not r.json()["committed"]
        )
        assert school.get("/api/students").json() == before
        r = school.post(
            "/api/students/import",
            headers=HEADERS,
            files=files,
            data={"confirm": "true"},
        )
        assert r.json()["committed"]
        ws.cell(2, 1, "Different identity")
        buf = BytesIO()
        wb.save(buf)
        r = school.post(
            "/api/students/import",
            headers=HEADERS,
            files={"file": ("students.xlsx", buf.getvalue())},
            data={"confirm": "true"},
        )
        assert r.json()["errors"] and not r.json()["committed"]
        assert school.get("/api/students").json()[-1]["name"] == "Aditi"
        assert (
            school.post("/api/classes", headers=HEADERS, json={"std": "11"}).status_code
            == 400
        )


def test_family_reports_and_reference_isolation(client):
    with TestClient(app) as parent, TestClient(app) as other:
        signup(parent, "PARENT")
        signup(other, "PARENT")
        profile = {
            "name": "Parent",
            "phone": "1234567890",
            "child_name": "Child",
            "child_class": "10",
        }
        for browser in (parent, other):
            assert (
                browser.put(
                    "/api/auth/profile", headers=HEADERS, json=profile
                ).status_code
                == 200
            )
        cls = parent.get("/api/classes").json()[0]
        subject = parent.post(
            "/api/subjects", headers=HEADERS, json={"name": "English"}
        ).json()
        e = exam(parent, cls, subject)
        child = parent.get("/api/students").json()[0]
        with SessionLocal() as db:
            sheet = AnswerSheetDocument(
                exam_id=e["id"],
                student_id=child["id"],
                original_filename="test.pdf",
                original_file="test.pdf",
                mime_type="application/pdf",
                file_size=1,
            )
            db.add(sheet)
            db.flush()
            ev = Evaluation(
                exam_id=e["id"],
                student_id=child["id"],
                answer_sheet_id=sheet.id,
                status="FINALIZED",
                total_marks=60,
                maximum_marks=80,
                percentage=75,
                finalized_at=datetime.now(timezone.utc),
            )
            db.add(ev)
            db.commit()
            ev_id = ev.id
        from app.services.school_reports import (
            build_evaluation_pdf_report,
            build_exam_pdf_report,
            build_report_cards,
        )

        with SessionLocal() as db:
            user = (
                db.query(Teacher)
                .filter_by(
                    workspace_id=e.get(
                        "workspace_id",
                        parent.get("/api/auth/me").json()["workspace_id"],
                    ),
                    role="PARENT",
                )
                .one()
            )
            for content, _ in (
                build_evaluation_pdf_report(db, user, ev_id),
                build_exam_pdf_report(db, user, e["id"]),
                build_report_cards(db, user, cls["id"], e["academic_year"], "Term 1"),
            ):
                pdf = fitz.open(stream=content, filetype="pdf")
                text = "".join(p.get_text() for p in pdf)
                assert "Eval AI" in text and "Pass" in text
                assert (
                    "A-1" not in text and "Rank" not in text and "Principal" not in text
                )
        assert other.get(f"/api/evaluations/{ev_id}").status_code == 404
        assert (
            other.get(f"/api/documents/answer-sheet/{sheet.id}/pdf").status_code == 404
        )
        ref = parent.post(
            "/api/reference-collection",
            headers=HEADERS,
            json={
                "source_document": "Private",
                "corrected_text": "Family private answer",
            },
        )
        assert ref.status_code == 201
        assert other.get("/api/reference-collection").json()["examples"] == []
        book = load_workbook(BytesIO(parent.get("/api/reports/marksheet").content))
        assert "Grade" not in [c.value for c in book.active[1]]
