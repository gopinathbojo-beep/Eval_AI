from io import BytesIO
from openpyxl import Workbook, load_workbook
from app.formatter.common.k_level import suggest_k_level
from .conftest import BEHAVIOUR, make_pdf, make_png


def j(r, code=None):
    if code is not None:
        assert r.status_code == code, f"{r.status_code}: {r.text}"
    return r.json()


def _xlsx(rows) -> bytes:
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _exam_ready_for_sheets(c, h, class_name):
    """Question paper + answer key uploaded together; the key waits and starts by itself once questions are verified."""
    cls = j(c.post("/api/classes", json={"name": class_name}, headers=h), 201)
    sub = j(c.post("/api/subjects", json={"name": f"Biology {class_name}"}, headers=h), 201)
    eid = j(c.post("/api/exams", headers=h, json={"name": "Bulk", "academic_year": "2026-27", "class_id": cls["id"],
                                                   "subject_id": sub["id"], "exam_type": "Unit", "total_marks": 15}), 201)["id"]
    key = j(c.post(f"/api/exams/{eid}/answer-key", headers=h, files={"file": ("key.png", make_png(), "image/png")}), 202)
    assert key["status"] == "WAITING"
    j(c.post(f"/api/exams/{eid}/question-paper", headers=h, files={"file": ("paper.pdf", make_pdf(2), "application/pdf")}), 202)
    assert j(c.get(f"/api/exams/{eid}/answer-key", headers=h))["status"] == "WAITING"   # untouched while questions change
    j(c.post(f"/api/exams/{eid}/questions/accept-valid", headers=h), 200)
    assert j(c.post(f"/api/exams/{eid}/questions/verify", headers=h), 200)["answer_key_started"] is True
    assert j(c.get(f"/api/exams/{eid}/answer-key", headers=h))["status"] == "READY"
    ks = j(c.get(f"/api/exams/{eid}/answer-keys", headers=h))["keys"]
    assert len(ks) == 4 and all(k["exam_question_id"] for k in ks)          # mapped onto the verified questions
    k2 = next(k for k in ks if k["raw_question_number"] == "2")
    j(c.patch(f"/api/answer-keys/{k2['id']}", json={"mismatch_resolution": "keep_question_marks"}, headers=h), 200)
    j(c.post(f"/api/exams/{eid}/answer-keys/verify", headers=h), 200)
    return cls, eid


def test_bulk_sheets_are_matched_to_students_or_held(client, auth):
    c, h = client, auth
    cls, eid = _exam_ready_for_sheets(c, h, "10-BulkA")
    mk = lambda name, roll, reg=None, section="A": j(c.post("/api/students", headers=h, json={  # noqa: E731
        "name": name, "roll_number": roll, "class_id": cls["id"], "registration_number": reg, "section": section}), 201)
    asha, ravi, meena, kumar = mk("Asha Devi", "1", "REG-001"), mk("Ravi Kumar", "2"), mk("Meena S", "3", "REG-003"), mk("Asha K", "4", "REG-004")

    BEHAVIOUR["sheet"] = "main"
    BEHAVIOUR["identities"] = [
        {"name": "Asha", "roll_number": "01", "registration_number": "reg 001"},     # reg + roll agree -> auto
        {"name": "Ravi Kumar", "roll_number": "2", "section": "A"},                  # roll only, no reg on file -> auto
        {"name": "Priya", "roll_number": "3"},                                      # roll 3 is Meena: name contradicts -> hold
        {}, {},                                                                     # nothing legible (header, then full page) -> hold
    ]
    files = [("files", (f"s{i}.pdf", make_pdf(1), "application/pdf")) for i in range(4)]
    up = j(c.post(f"/api/exams/{eid}/answer-sheets/bulk", headers=h, files=files), 202)
    assert up["uploaded"] == 4

    sheets = j(c.get(f"/api/exams/{eid}/answer-sheets", headers=h))
    assert {s["student"]["id"] for s in sheets} == {asha["id"], ravi["id"]}
    assert all(s["detected_answers"] > 0 for s in sheets)                       # the regular pipeline ran on them
    held = j(c.get(f"/api/exams/{eid}/answer-sheets/unassigned", headers=h))
    assert [d["status"] for d in held] == ["NEEDS_STUDENT", "NEEDS_STUDENT"]
    priya, blank = held
    # held sheets are read anyway, so the batch never waits on the teacher; only the student is left to pick
    assert all(pg["ocr_status"] == "OK" for d in held for pg in d["pages"])
    assert {d["stage"] for d in held} == {"Pick the student (answers already read)"}
    assert "Priya" in priya["identity"]["note"] and priya["candidates"][0]["id"] == meena["id"]
    assert "No roll number or registration number" in blank["identity"]["note"]

    # teacher resolves: Priya's sheet is Meena's; the blank one cannot go to Asha (she already has one)
    BEHAVIOUR["sheet"] = "dup"                  # a different sheet: if assigning read the pages again, its answers would show
    r = j(c.post(f"/api/answer-sheets/{priya['id']}/assign", headers=h, json={"student_id": meena["id"]}), 200)
    BEHAVIOUR["sheet"] = "main"
    assert r["student"]["id"] == meena["id"]
    answers = j(c.get(f"/api/evaluations/{r['id']}", headers=h))["answers"]
    assert answers and not any("Another try" in (a["answer_text"] or "") for a in answers)
    assert c.post(f"/api/answer-sheets/{blank['id']}/assign", headers=h, json={"student_id": asha["id"]}).status_code == 409
    # a sheet whose student already has one is held, never auto-assigned twice
    BEHAVIOUR["identities"] = [{"registration_number": "REG-004", "roll_number": "1"}]   # reg -> Kumar, roll -> Asha
    j(c.post(f"/api/answer-sheets/{blank['id']}/identify", headers=h), 202)
    again = next(d for d in j(c.get(f"/api/exams/{eid}/answer-sheets/unassigned", headers=h)) if d["id"] == blank["id"])
    assert again["status"] == "NEEDS_STUDENT" and "roll number 1" in again["identity"]["note"]
    j(c.post(f"/api/answer-sheets/{blank['id']}/assign", headers=h, json={"student_id": kumar["id"]}), 200)
    assert j(c.get(f"/api/exams/{eid}/answer-sheets/unassigned", headers=h)) == []
    assert len(j(c.get(f"/api/exams/{eid}/answer-sheets", headers=h))) == 4

    # an unreadable upload in a batch rejects the batch before anything is stored
    bad = [("files", ("ok.pdf", make_pdf(1), "application/pdf")), ("files", ("bad.pdf", b"nope", "application/pdf"))]
    r = c.post(f"/api/exams/{eid}/answer-sheets/bulk", headers=h, files=bad)
    assert r.status_code == 400 and "bad.pdf" in r.json()["message"]
    assert j(c.get(f"/api/exams/{eid}/answer-sheets/unassigned", headers=h)) == []


def test_bulk_sheets_are_graded_in_turn_and_finalized_together(client, auth):
    c, h = client, auth
    cls, eid = _exam_ready_for_sheets(c, h, "10-BulkB")
    mk = lambda name, roll: j(c.post("/api/students", headers=h, json={  # noqa: E731
        "name": name, "roll_number": roll, "class_id": cls["id"]}), 201)
    clean, messy, unsure = mk("Clean Kid", "1"), mk("Messy Kid", "2"), mk("Unsure Kid", "3")

    for sheet, roll in (("clean", "1"), ("main", "2"), ("unsure", "3")):
        BEHAVIOUR["sheet"], BEHAVIOUR["identities"] = sheet, [{"roll_number": roll}]
        j(c.post(f"/api/exams/{eid}/answer-sheets/bulk", headers=h,
                 files=[("files", (f"{sheet}.pdf", make_pdf(1), "application/pdf"))]), 202)
    BEHAVIOUR["sheet"] = "main"

    # every sheet was read and the clean ones graded without the teacher starting anything
    sheets = {s["student"]["id"]: s for s in j(c.get(f"/api/exams/{eid}/answer-sheets", headers=h))}
    assert sheets[clean["id"]]["run_status"] == "DONE" and sheets[unsure["id"]]["run_status"] == "DONE"
    assert sheets[messy["id"]]["run_status"] == "IDLE"                          # unmapped answer: left for the teacher

    summ = j(c.get(f"/api/exams/{eid}/review-summary", headers=h), 200)
    assert summ["processing"] == 0 and summ["total"] == 3 and summ["unassigned"] == 0
    assert [r["student"]["id"] for r in summ["ready"]] == [clean["id"]]
    review = {r["student"]["id"]: r["reasons"] for r in summ["needs_review"]}
    assert set(review) == {messy["id"], unsure["id"]}
    assert any("Fix the answer mapping" in x["reason"] for x in review[messy["id"]])
    assert not any(x["reason"] == "Not evaluated yet" for x in review[messy["id"]])
    assert [x["question"] for x in review[unsure["id"]]] == ["3(b)"] and "30% confidence" in review[unsure["id"]][0]["reason"]

    done = j(c.post(f"/api/exams/{eid}/finalize-ready", headers=h), 200)
    assert done["finalized_now"] == 1 and done["finalized"] == 1 and done["ready"] == [] and len(done["needs_review"]) == 2
    res = j(c.get(f"/api/evaluations/{sheets[clean['id']]['id']}/result", headers=h))
    assert res["finalized"] and res["total_marks"] == res["maximum_marks"]
    r = c.post(f"/api/exams/{eid}/finalize-ready", headers=h)
    assert r.status_code == 400 and r.json()["code"] == "NOTHING_READY"

    # once the teacher checks the unsure mark, that sheet joins the next bulk finalize
    ev = j(c.get(f"/api/evaluations/{sheets[unsure['id']]['id']}", headers=h))
    it = next(i for i in ev["items"] if i["ai_confidence"] is not None and i["ai_confidence"] < 0.5)
    j(c.patch(f"/api/evaluations/{ev['evaluation']['id']}/items/{it['id']}", headers=h, json={"final_marks": 2, "approve": True, "reason": "Checked"}), 200)
    assert j(c.post(f"/api/exams/{eid}/finalize-ready", headers=h), 200)["finalized_now"] == 1


def test_student_import_export(client, auth):
    c, h = client, auth
    data = _xlsx([["Class list"], ["Student Name", "Roll No", "Reg No", "Class", "Sec"],
                  ["Lakshmi", 1, "IMP-1", "12", "ImportA"], ["Arun", 2, "IMP-2", "12", "ImportB"], ["Bala", 3, "IMP-3", "12", "ImportA"],
                  ["Dup Roll", 1, None, "12", "ImportA"], ["", "", "", "", ""], ["No Roll", None, None, "12", None]])
    r = j(c.post("/api/students/import", headers=h, files={"file": ("s.xlsx", data, "application/octet-stream")}), 200)
    assert r["created"] == 3 and r["updated"] == 0 and [e["row"] for e in r["errors"]] == [6, 8]
    assert not r["committed"]
    data = _xlsx([["Name", "Roll Number", "Registration Number", "Class", "Section"],
                  ["Lakshmi", 1, "IMP-1", "12", "ImportA"], ["Arun", 2, "IMP-2", "12", "ImportB"], ["Bala", 3, "IMP-3", "12", "ImportA"]])
    assert j(c.post("/api/students/import", headers=h, files={"file": ("s.xlsx", data)}, data={"confirm":"true"}), 200)["committed"]
    names = {x["name"] for x in j(c.get("/api/classes", headers=h))}
    assert {"12-ImportA", "12-ImportB"} <= names
    # re-import updates by registration number instead of duplicating
    data = _xlsx([["Name", "Roll Number", "Registration Number", "Class", "Section"], ["Lakshmi", 1, "IMP-1", "12", "ImportA"]])
    assert j(c.post("/api/students/import", headers=h, files={"file": ("s.xlsx", data, "x")}), 200)["updated"] == 1
    assert c.post("/api/students/import", headers=h, files={"file": ("s.csv", b"a,b", "text/csv")}).status_code == 400

    cls = next(x for x in j(c.get("/api/classes", headers=h)) if x["name"] == "12-ImportA")
    st = j(c.get(f"/api/students?class_id={cls['id']}", headers=h))
    assert {(s["name"], s["registration_number"], s["section"]) for s in st} == {("Lakshmi", "IMP-1", "ImportA"), ("Bala", "IMP-3", "ImportA")}
    r = c.post("/api/students", headers=h, json={"name": "X", "roll_number": "9", "class_id": cls["id"], "registration_number": "imp 1"})
    assert r.status_code == 409 and r.json()["code"] == "DUPLICATE_REG"

    x = c.get(f"/api/students/export?format=xlsx&class_id={cls['id']}", headers=h)
    assert x.status_code == 200
    rows = list(load_workbook(BytesIO(x.content)).active.iter_rows(values_only=True))
    assert rows[0] == ("Name", "Roll Number", "Registration Number", "Class", "Section") and len(rows) == 3
    assert rows[1][3] == "12-ImportA"
    p = c.get("/api/students/export?format=pdf", headers=h)
    assert p.status_code == 200 and p.content.startswith(b"%PDF")


def test_k_level_suggestions():
    assert suggest_k_level("Define osmosis.") == "K1"
    assert suggest_k_level("Explain diffusion with an example.") == "K2"
    assert suggest_k_level("Solve for x: 2x + 3 = 7") == "K3"
    assert suggest_k_level("Evaluate (3/4)² × 8") == "K3"
    assert suggest_k_level("Compare mitosis and meiosis.") == "K4"
    assert suggest_k_level("Evaluate the impact of deforestation and justify your view.") == "K5"
    assert suggest_k_level("Design an experiment to show transpiration.") == "K6"
    assert suggest_k_level("Write a letter to the editor about pollution.") == "K6"
    assert suggest_k_level("") is None


def test_k_level_suggestion_reads_tamil_command_words():
    """Tamil papers have no English verbs; without the Tamil words every question fell back to K1."""
    from app import formatter
    with formatter.use("tamil"):
        assert suggest_k_level("அ) கல்வியின் சிறப்புகள் யாவை? (2)") == "K1"
        assert suggest_k_level("அ) கீழ்க்கண்ட சொற்களின் எதிர்ச்சொற்களை எழுதுக. ஆ) பன்மை வடிவங்களை எழுதுக.") == "K1"
        assert suggest_k_level("ஆ) இயற்கையை நாம் ஏன் பாதுகாக்க வேண்டும்?") == "K2"
        assert suggest_k_level("படித்துப் புரிந்து விடையளித்தல் (3)") == "K2"
        assert suggest_k_level("கீழ்க்காணும் சொற்களுக்குப் பன்மை வடிவம் எழுதுக.") == "K3"
        assert suggest_k_level("கடிதம் எழுதுதல் (5) / நண்பருக்குக் கடிதம் எழுதுக.") == "K6"
        assert suggest_k_level("கட்டுரை எழுதுதல் (5)") == "K6"
        assert suggest_k_level("அவன் தயார் நிலையில் இருந்தான்.") == "K1"  # "யார்" inside "தயார்" is no question word
    with formatter.use("english"):
        assert suggest_k_level("கட்டுரை எழுதுதல்") == "K1"  # Tamil words only under a Tamil subject
