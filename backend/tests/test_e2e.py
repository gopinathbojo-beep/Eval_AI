import pytest
from .conftest import BEHAVIOUR, MODEL, make_pdf, make_png

H = {}


def j(r, code=None):
    if code is not None:
        assert r.status_code == code, f"{r.status_code}: {r.text}"
    return r.json()


def test_auth_required(client):
    assert client.get("/api/exams").status_code == 401
    assert client.get("/api/dashboard").status_code == 401
    assert client.post("/api/auth/login", json={"email": "rao@school.org", "password": "nope"}).status_code == 401
    r = client.get("/api/exams", headers={"Authorization": "Bearer garbage"})
    assert r.status_code == 401 and r.json()["code"] == "UNAUTHENTICATED"


def test_ai_status(client, auth):
    s = j(client.get("/api/system/ai-status", headers=auth), 200)
    assert s["available"] and s["model"] == MODEL


def test_full_workflow(client, auth):
    c, h = client, auth
    cls = j(c.post("/api/classes", json={"name": "10-A"}, headers=h), 201)
    sub = j(c.post("/api/subjects", json={"name": "Biology"}, headers=h), 201)
    exam = j(c.post("/api/exams", headers=h, json={"name": "Term 1", "academic_year": "2026-27", "class_id": cls["id"],
             "subject_id": sub["id"], "exam_type": "Unit", "total_marks": 15}), 201)
    eid = exam["id"]
    assert exam["status"] == "DRAFT"

    # --- bad uploads are rejected
    r = c.post(f"/api/exams/{eid}/question-paper", headers=h, files={"file": ("x.pdf", b"not a pdf", "application/pdf")})
    assert r.status_code == 400
    r = c.post(f"/api/exams/{eid}/question-paper", headers=h, files={"file": ("../../evil.pdf", make_pdf(), "application/pdf")})
    assert r.status_code == 400
    r = c.post(f"/api/exams/{eid}/question-paper", headers=h, files={"file": ("x.exe", b"MZ", "application/octet-stream")})
    assert r.status_code == 400
    # (an answer key uploaded before the questions are verified waits for them: see test_bulk_sheets_and_early_key.py)

    # --- question paper (page 2 fails to parse first)
    BEHAVIOUR["garbage_pages"] = {2}
    j(c.post(f"/api/exams/{eid}/question-paper", headers=h, files={"file": ("paper.pdf", make_pdf(2), "application/pdf")}), 202)
    qp = j(c.get(f"/api/exams/{eid}/question-paper", headers=h), 200)
    assert qp["status"] == "NEEDS_ATTENTION"
    assert [p["ocr_status"] for p in qp["pages"]] == ["OK", "FAILED_PARSE"]
    diag = j(c.get(f"/api/documents/question-paper/{qp['id']}/pages/2/ocr", headers=h), 200)
    assert diag["status"] == "FAILED_PARSE" and "cannot read" in diag["raw"]
    assert "InternVL3.5 4B failed to return a valid structured response" in qp["error_message"]
    # verification is blocked while a page failed
    r = c.post(f"/api/exams/{eid}/questions/verify", headers=h)
    assert r.status_code == 400 and r.json()["code"] == "QUESTIONS_NOT_READY"
    # retry only the failed page
    BEHAVIOUR["garbage_pages"] = set()
    j(c.post(f"/api/documents/question-paper/{qp['id']}/reprocess", headers=h), 202)
    qp = j(c.get(f"/api/exams/{eid}/question-paper", headers=h), 200)
    assert qp["status"] == "READY"
    qs = j(c.get(f"/api/exams/{eid}/questions", headers=h), 200)
    nums = [q["display_number"] for q in qs["questions"]]
    assert nums == ["1", "2", "3(a)", "3(b)"]
    assert all(q["normalized_number"] == q["display_number"] for q in qs["questions"])
    assert [q["max_marks"] for q in qs["questions"]] == [2, 5, 3, 5]
    assert qs["questions"][0]["has_crop"]
    assert j(c.get(f"/api/exams/{eid}", headers=h))["status"] == "QUESTIONS_EXTRACTED"

    # verify requires explicit acceptance
    r = c.post(f"/api/exams/{eid}/questions/verify", headers=h)
    assert r.status_code == 400 and len(r.json()["issues"]) == 4
    # edit: bad number rejected, marks edit audited
    q1 = qs["questions"][0]
    assert c.patch(f"/api/exam-questions/{q1['id']}", json={"display_number": "banana"}, headers=h).status_code == 400
    j(c.patch(f"/api/exam-questions/{q1['id']}", json={"max_marks": 2.0}, headers=h), 200)
    j(c.post(f"/api/exams/{eid}/questions/accept-valid", headers=h), 200)
    v = j(c.post(f"/api/exams/{eid}/questions/verify", headers=h), 200)
    assert v["verified"] == 4 and v["marks_warning"] is None
    assert j(c.get(f"/api/exams/{eid}", headers=h))["status"] == "QUESTIONS_VERIFIED"
    # locked after verification
    assert c.patch(f"/api/exam-questions/{q1['id']}", json={"max_marks": 9}, headers=h).status_code == 409
    assert c.post(f"/api/exams/{eid}/question-paper", headers=h, files={"file": ("p.pdf", make_pdf(), "application/pdf")}).status_code == 409
    qid = {q["display_number"]: q["id"] for q in qs["questions"]}

    # --- answer key: labels "Q1", "2", "3 (a)", "3b" must map deterministically; Q2 marks disagree (4 vs 5)
    j(c.post(f"/api/exams/{eid}/answer-key", headers=h, files={"file": ("key.png", make_png(), "image/png")}), 202)
    ks = j(c.get(f"/api/exams/{eid}/answer-keys", headers=h), 200)
    assert len(ks["keys"]) == 4 and not ks["missing_question_ids"]
    by_label = {k["raw_question_number"]: k for k in ks["keys"]}
    assert by_label["Q1"]["exam_question_id"] == qid["1"]
    assert by_label["2"]["exam_question_id"] == qid["2"]
    assert by_label["3 (a)"]["exam_question_id"] == qid["3(a)"]
    assert by_label["3b"]["exam_question_id"] == qid["3(b)"]
    assert by_label["2"]["mark_mismatch"] is True
    r = c.post(f"/api/exams/{eid}/answer-keys/verify", headers=h)
    assert r.status_code == 400 and any("MARK MISMATCH" in i["reason"] for i in r.json()["issues"])
    j(c.patch(f"/api/answer-keys/{by_label['2']['id']}", json={"mismatch_resolution": "keep_question_marks"}, headers=h), 200)
    assert j(c.post(f"/api/exams/{eid}/answer-keys/verify", headers=h), 200)["verified"] == 4
    assert j(c.get(f"/api/exams/{eid}", headers=h))["status"] == "ANSWER_KEY_VERIFIED"
    # question's marks never silently changed
    qs2 = j(c.get(f"/api/exams/{eid}/questions", headers=h))
    assert [q["max_marks"] for q in qs2["questions"]] == [2, 5, 3, 5]

    # --- student answer sheet
    BEHAVIOUR["sheet"] = "main"
    r = c.post(f"/api/exams/{eid}/answer-sheets", headers=h, data={"student_name": "Asha", "roll_number": "7"},
               files={"file": ("sheet.pdf", make_pdf(1), "application/pdf")})
    ev = j(r, 202)
    evid = ev["id"]
    assert c.post(f"/api/exams/{eid}/answer-sheets", headers=h, data={"student_name": "Asha", "roll_number": "7"},
                  files={"file": ("again.pdf", make_pdf(1), "application/pdf")}).status_code == 409
    d = j(c.get(f"/api/evaluations/{evid}", headers=h), 200)
    st = {a["raw_question_number"]: a for a in d["answers"]}
    assert st["Q1"]["status"] == "MAPPED" and st["Q1"]["exam_question_id"] == qid["1"]
    assert st["Q2"]["exam_question_id"] == qid["2"]
    assert st["3(b)"]["exam_question_id"] == qid["3(b)"]
    assert st["3"]["status"] == "NEEDS_REVIEW" and st["3"]["exam_question_id"] is None          # never guessed
    assert set(st["3"]["mapping_candidates"]) == {qid["3(a)"], qid["3(b)"]}
    assert st[None]["status"] == "UNMAPPED"
    assert d["evaluation"]["status"] == "MAPPING_REVIEW"
    assert d["summary"]["needs_review"] == 1 and d["summary"]["unmapped"] == 1 and d["summary"]["missing"] == 1
    assert j(c.get(f"/api/exams/{eid}", headers=h))["status"] == "READY_FOR_EVALUATION"

    # evaluation + finalize are refused with structured reasons
    r = c.post(f"/api/evaluations/{evid}/run", headers=h)
    assert r.status_code == 400 and r.json()["code"] == "MAPPING_INCOMPLETE" and len(r.json()["issues"]) == 2
    r = c.post(f"/api/evaluations/{evid}/finalize", headers=h)
    assert r.status_code == 400 and r.json()["code"] == "EVALUATION_NOT_READY" and r.json()["issues"]

    # teacher resolves mapping: 3 -> 3(a); stray text ignored
    j(c.patch(f"/api/evaluations/{evid}/mapping/{st['3']['id']}", json={"exam_question_id": qid["3(a)"]}, headers=h), 200)
    j(c.patch(f"/api/evaluations/{evid}/mapping/{st[None]['id']}", json={"ignore": True}, headers=h), 200)
    d = j(c.get(f"/api/evaluations/{evid}", headers=h))
    assert d["summary"]["missing"] == 0 and d["evaluation"]["status"] == "OCR_REVIEW"

    # --- AI evaluation
    j(c.post(f"/api/evaluations/{evid}/run", headers=h), 202)
    d = j(c.get(f"/api/evaluations/{evid}", headers=h))
    it = {i["exam_question_id"]: i for i in d["items"]}
    assert d["evaluation"]["run_status"] == "DONE"
    assert it[qid["1"]]["ai_marks"] == 2 and it[qid["2"]]["ai_marks"] == 0
    assert it[qid["3(b)"]]["status"] == "EVALUATION_ERROR" and it[qid["3(b)"]]["ai_marks"] is None   # AI said 8 of 5
    assert "exceeds" in it[qid["3(b)"]]["error"]
    # finalize still blocked; issue names the question
    r = c.post(f"/api/evaluations/{evid}/finalize", headers=h)
    assert r.status_code == 400 and any(i["question"] == "3(b)" for i in r.json()["issues"])
    # approving an errored item without marks is refused
    r = c.patch(f"/api/evaluations/{evid}/items/{it[qid['3(b)']]['id']}", json={"approve": True}, headers=h)
    assert r.status_code == 400
    # teacher fixes the OCR text -> evaluation of just that question re-runs
    a3b = st["3(b)"]["id"]
    j(c.patch(f"/api/evaluations/{evid}/mapping/{a3b}", json={"answer_text": "Leaf diagram with labels", "reason": "OCR typo"}, headers=h), 200)
    d = j(c.get(f"/api/evaluations/{evid}", headers=h))
    it = {i["exam_question_id"]: i for i in d["items"]}
    assert it[qid["3(b)"]]["status"] == "EVALUATED" and it[qid["3(b)"]]["ai_marks"] == 5
    ans3b = next(a for a in d["answers"] if a["id"] == a3b)
    assert "OVERSHOOT" in ans3b["raw_ocr_text"]                       # raw OCR never overwritten
    assert ans3b["answer_text"] == "Leaf diagram with labels" and ans3b["edited"]
    # override requires a reason, and the AI mark is preserved
    q2item = it[qid["2"]]
    assert c.patch(f"/api/evaluations/{evid}/items/{q2item['id']}", json={"final_marks": 2}, headers=h).status_code == 400
    assert c.patch(f"/api/evaluations/{evid}/items/{q2item['id']}", json={"final_marks": 9, "reason": "x"}, headers=h).status_code == 400
    r = j(c.patch(f"/api/evaluations/{evid}/items/{q2item['id']}", json={"final_marks": 1, "reason": "Partial credit for definition"}, headers=h), 200)
    assert r["item"]["final_marks"] == 1 and r["item"]["ai_marks"] == 0 and r["item"]["teacher_overridden"]
    # items still awaiting approval -> not finalizable; then approve all
    assert c.post(f"/api/evaluations/{evid}/finalize", headers=h).status_code == 400
    j(c.patch(f"/api/evaluations/{evid}/items/{q2item['id']}", json={"approve": True}, headers=h), 200)
    j(c.post(f"/api/evaluations/{evid}/items/approve-all", headers=h), 200)
    d = j(c.get(f"/api/evaluations/{evid}", headers=h))
    assert d["evaluation"]["status"] == "FINALIZABLE" and d["issues"] == []
    fin = j(c.post(f"/api/evaluations/{evid}/finalize", headers=h), 200)
    assert fin["status"] == "FINALIZED"
    assert (fin["total_marks"], fin["maximum_marks"], fin["percentage"]) == (2 + 1 + 3 + 5, 15, 73.33)
    # frozen afterwards
    assert c.patch(f"/api/evaluations/{evid}/mapping/{a3b}", json={"answer_text": "z"}, headers=h).status_code == 409

    res = j(c.get(f"/api/evaluations/{evid}/result", headers=h), 200)
    assert res["finalized"] and [(q["question"], q["final_marks"], q["max_marks"]) for q in res["questions"]] == \
        [("1", 2, 2), ("2", 1, 5), ("3(a)", 3, 3), ("3(b)", 5, 5)]
    assert res["questions"][1]["ai_marks"] == 0 and res["questions"][1]["official_answer"]
    assert "k_level" not in res["questions"][0] and "k_levels" not in res["outcomes"]
    j(c.patch(f"/api/exam-questions/{qid['2']}/outcome", json={"course_outcome": "2"}, headers=h), 200)
    summ = j(c.get(f"/api/exams/{eid}/outcome-summary", headers=h), 200)
    assert summ["finalized_students"] == 1 and "k_levels" not in summ
    assert [r["key"] for r in summ["course_outcomes"]] == ["CO2", "Untagged"]
    rep_pdf = c.get(f"/api/evaluations/{evid}/report.pdf", headers=h)
    assert rep_pdf.status_code == 200 and rep_pdf.content.startswith(b"%PDF")
    from io import BytesIO
    from openpyxl import load_workbook
<<<<<<< 7c662e7
    assert "K-Level Analysis" not in load_workbook(BytesIO(c.get("/api/reports/marksheet", headers=h).content)).sheetnames
=======
    workbook = load_workbook(BytesIO(c.get("/api/reports/marksheet", headers=h).content))
    assert "K-Level Analysis" not in workbook.sheetnames
    assert "Grade" in [cell.value for cell in workbook.active[1]]
>>>>>>> kishore/Version2.0
    # audit trail
    logs = j(c.get(f"/api/exams/{eid}/audit-logs", headers=h), 200)
    actions = {l["action"] for l in logs}
    assert {"marks_changed", "student_answer_mapping_changed", "evaluation_finalized", "question_max_marks_changed"} & actions
    assert "answer_key_marks_aligned_to_question" in actions and "questions_verified" in actions
    # exam finalize: every student's report is saved, files are deleted 24 h later (or on demand)
    assert c.post(f"/api/exams/{eid}/purge-files", headers=h).status_code == 409        # not finalized yet
    fin = j(c.post(f"/api/exams/{eid}/finalize", headers=h), 200)
    assert fin["status"] == "FINALIZED" and fin["reports_saved"] == 1 and fin["files_purge_at"]
    from datetime import datetime, timedelta, timezone
    from app.database import SessionLocal
    from app.models import Exam
    from app.services.exam_files import exam_dir, purge_due
    assert purge_due() == 0                                                           # still inside the grace period
    root = exam_dir(eid)
    assert {p.name for p in root.iterdir()} >= {"question-paper", "answer-key", "answer-sheets", "reports"}
    db = SessionLocal()
    db.get(Exam, eid).files_purge_at = datetime.now(timezone.utc) - timedelta(minutes=1)
    db.commit()
    db.close()
    assert purge_due() == 1
    assert [p.name for p in root.iterdir()] == ["reports"] and len(list((root / "reports").iterdir())) == 1
    assert j(c.get(f"/api/exams/{eid}", headers=h))["files_purged_at"]
    r = c.get(f"/api/documents/answer-sheet/{d['document']['id']}/pages/1/image", headers=h)
    assert r.status_code == 404 and r.json()["code"] == "FILES_DELETED"
    # results and the saved official report still work without the files
    assert j(c.get(f"/api/evaluations/{evid}/result", headers=h), 200)["total_marks"] == 11
    rep_pdf = c.get(f"/api/evaluations/{evid}/report.pdf", headers=h)
    assert rep_pdf.status_code == 200 and rep_pdf.content.startswith(b"%PDF")
    assert "attachment" in rep_pdf.headers["content-disposition"] and "__" not in rep_pdf.headers["content-disposition"]


def test_duplicate_mapping_and_ai_unavailable(client, auth):
    c, h = client, auth
    cls = j(c.post("/api/classes", json={"name": "12-B"}, headers=h), 201)
    sub = j(c.post("/api/subjects", json={"name": "Botany"}, headers=h), 201)
    eid = j(c.post("/api/exams", headers=h, json={"name": "T2", "academic_year": "26", "class_id": cls["id"],
            "subject_id": sub["id"], "exam_type": "Unit", "total_marks": 15}), 201)["id"]
    # Ollama down -> clear AI_ERROR, nothing invented
    from app.config import get_settings
    s = get_settings(); good = s.internvl_api_url
    s.internvl_api_url = "http://127.0.0.1:1"
    try:
        j(c.post(f"/api/exams/{eid}/question-paper", headers=h, files={"file": ("p.pdf", make_pdf(2), "application/pdf")}), 202)
        qp = j(c.get(f"/api/exams/{eid}/question-paper", headers=h))
        assert qp["status"] == "AI_ERROR"
        assert qp["error_message"] == "InternVL3.5 4B is not available. Start Ollama and ensure blaifa/InternVL3_5:4B is installed."
        assert j(c.get(f"/api/exams/{eid}/questions", headers=h))["questions"] == []
        st = j(c.get("/api/system/ai-status", headers=h))
        assert st["available"] is False
    finally:
        s.internvl_api_url = good
    j(c.post(f"/api/documents/question-paper/{qp['id']}/reprocess", headers=h), 202)
    assert j(c.get(f"/api/exams/{eid}/question-paper", headers=h))["status"] == "READY"
    j(c.post(f"/api/exams/{eid}/questions/accept-valid", headers=h), 200)
    j(c.post(f"/api/exams/{eid}/questions/verify", headers=h), 200)
    j(c.post(f"/api/exams/{eid}/answer-key", headers=h, files={"file": ("k.png", make_png(), "image/png")}), 202)
    ks = j(c.get(f"/api/exams/{eid}/answer-keys", headers=h))["keys"]
    j(c.patch(f"/api/answer-keys/{next(k for k in ks if k['raw_question_number']=='2')['id']}", json={"mismatch_resolution": "keep_question_marks"}, headers=h))
    j(c.post(f"/api/exams/{eid}/answer-keys/verify", headers=h), 200)

    BEHAVIOUR["sheet"] = "dup"
    ev = j(c.post(f"/api/exams/{eid}/answer-sheets", headers=h, data={"student_name": "Ravi", "roll_number": "3"},
                  files={"file": ("s.png", make_png(), "image/png")}), 202)
    d = j(c.get(f"/api/evaluations/{ev['id']}", headers=h))
    q1 = [a for a in d["answers"] if a["exam_question_id"] and a["normalized_number"] == "1"]
    assert len(q1) == 2 and all(a["status"] == "NEEDS_REVIEW" and a["mapping_note"].startswith("DUPLICATE_MAPPING") for a in q1)
    r = c.post(f"/api/evaluations/{ev['id']}/run", headers=h)
    assert r.status_code == 400 and r.json()["code"] == "MAPPING_INCOMPLETE"
    # ignoring one duplicate resolves the other automatically
    j(c.patch(f"/api/evaluations/{ev['id']}/mapping/{q1[1]['id']}", json={"ignore": True}, headers=h), 200)
    d = j(c.get(f"/api/evaluations/{ev['id']}", headers=h))
    assert next(a for a in d["answers"] if a["id"] == q1[0]["id"])["status"] == "MAPPED"
    assert d["summary"]["needs_review"] == 0
    BEHAVIOUR["sheet"] = "main"
